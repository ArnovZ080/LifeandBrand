"""
Micros POS export adapter (Phase 1).

Reads daily sales export files (.csv / .xlsx / .xlsb) dropped into a watch
directory and turns them into canonical sales rows for the ingest pipeline.

Design note — duplicate detection:
The adapter is deliberately DB-free. The pipeline (which owns the session)
queries raw_ingests for known sha256 hashes and passes them into
``__init__(known_hashes=...)``. ``discover()`` then hashes each candidate
file and skips ones already ingested. This keeps the adapter unit-testable
without a database and keeps persistence concerns in the pipeline.
"""
from __future__ import annotations

import csv
import hashlib
import os
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from app.connectors.base import SourceAdapter

SUPPORTED_EXTENSIONS = (".csv", ".xlsx", ".xlsb")

# Canonical field → accepted (normalised) header names, in priority order.
HEADER_ALIASES: dict[str, list[str]] = {
    "pos_site_code": ["pos_site_code", "site_code", "store_code", "location_code",
                      "site", "store", "location", "rvc", "revenue_center", "revenue_centre"],
    "business_date": ["business_date", "business_day", "trans_date", "transaction_date",
                      "sale_date", "date", "day"],
    "gross_sales": ["gross_sales", "gross", "total_revenue", "total_sales",
                    "sales_total", "gross_total", "revenue"],
    "vat": ["vat", "vat_amount", "tax", "tax_total", "tax_amount"],
    "discounts": ["discounts", "discount", "discount_total", "discount_amount", "promo", "promotions"],
    "net_sales": ["net_sales", "net", "net_total", "net_amount", "net_revenue"],
    "covers": ["covers", "guests", "guest_count", "cover_count", "pax", "customers"],
    # Item-level (optional) columns
    "pos_item_code": ["pos_item_code", "item_code", "item_number", "menu_item_id",
                      "micros_item_id", "plu", "plu_code"],
    "item_name": ["item_name", "menu_item", "item_description", "description", "item"],
    "qty": ["qty", "quantity", "quantity_sold", "sold_qty", "sales_count", "count_sold"],
    "gross_value": ["gross_value", "line_total", "amount", "value", "item_revenue", "line_amount"],
}


def sha256_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _normalise_header(name: str) -> str:
    """strip / lower / collapse non-alphanumerics to underscores."""
    return re.sub(r"[^a-z0-9]+", "_", str(name).strip().lower()).strip("_")


def _to_decimal(value) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    s = str(value).strip().replace(",", "").replace(" ", "")
    s = re.sub(r"[^0-9.\-()]", "", s)  # strip currency symbols
    if not s:
        return None
    if s.startswith("(") and s.endswith(")"):  # accounting negatives
        s = "-" + s[1:-1]
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _to_iso_date(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        # Excel serial date (days since 1899-12-30)
        try:
            from datetime import timedelta
            return (date(1899, 12, 30) + timedelta(days=int(value))).isoformat()
        except (OverflowError, ValueError):
            return None
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d",
                "%d %b %Y", "%d %B %Y", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


class MicrosExportAdapter(SourceAdapter):
    """Adapter for file-drop Micros exports. See module docstring for design."""

    def __init__(self, known_hashes: set[str] | None = None):
        self.known_hashes: set[str] = known_hashes or set()

    def source_key(self) -> str:
        return "micros"

    # ── discover ────────────────────────────────────────────────────────────────
    def discover(self, watch_dir: str) -> list[str]:
        """Return paths of supported files in watch_dir whose sha256 is not
        in ``self.known_hashes``. Missing directory → empty list."""
        if not os.path.isdir(watch_dir):
            return []
        paths = []
        for name in sorted(os.listdir(watch_dir)):
            path = os.path.join(watch_dir, name)
            if not os.path.isfile(path):
                continue
            if not name.lower().endswith(SUPPORTED_EXTENSIONS):
                continue
            if sha256_file(path) in self.known_hashes:
                continue
            paths.append(path)
        return paths

    # ── parse ───────────────────────────────────────────────────────────────────
    def parse(self, path: str) -> dict:
        """Read the file into {file_name, rows: [{col: value}]} with
        normalised (strip/lower/underscore) column names."""
        ext = os.path.splitext(path)[1].lower()
        if ext == ".csv":
            rows = self._parse_csv(path)
        elif ext == ".xlsx":
            rows = self._parse_xlsx(path)
        elif ext == ".xlsb":
            rows = self._parse_xlsb(path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")
        return {"file_name": os.path.basename(path), "rows": rows}

    @staticmethod
    def _rows_from_matrix(matrix: list[list]) -> list[dict]:
        """First non-empty row is the header; remaining rows become dicts."""
        rows: list[dict] = []
        header: list[str] | None = None
        for raw in matrix:
            if raw is None or all(c is None or str(c).strip() == "" for c in raw):
                continue
            if header is None:
                header = [_normalise_header(c) for c in raw]
                continue
            row = {}
            for i, cell in enumerate(raw):
                if i < len(header) and header[i]:
                    row[header[i]] = cell
            if row:
                rows.append(row)
        return rows

    def _parse_csv(self, path: str) -> list[dict]:
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
            return self._rows_from_matrix(list(csv.reader(f)))

    def _parse_xlsx(self, path: str) -> list[dict]:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            return self._rows_from_matrix([list(r) for r in ws.iter_rows(values_only=True)])
        finally:
            wb.close()

    def _parse_xlsb(self, path: str) -> list[dict]:
        from pyxlsb import open_workbook
        with open_workbook(path) as wb:
            with wb.get_sheet(1) as sheet:
                return self._rows_from_matrix(
                    [[c.v for c in row] for row in sheet.rows()]
                )

    # ── normalise ───────────────────────────────────────────────────────────────
    def normalise(self, payload: dict) -> list[dict]:
        """Map raw rows to canonical dicts, one per (pos_site_code, business_date):

            { pos_site_code, business_date (ISO), gross_sales, vat, discounts,
              net_sales, covers, lines: [{pos_item_code, item_name, qty, gross_value}] }

        Rows carrying item-level columns are aggregated into the day's totals
        and appended to that day's ``lines``. Rows that cannot be mapped
        (missing site or date) are collected into ``payload["skipped_rows"]``
        instead of failing the whole file.
        """
        skipped: list[dict] = []
        buckets: dict[tuple[str, str], dict] = {}

        for raw in payload.get("rows", []):
            mapped = {field: self._pick(raw, field) for field in HEADER_ALIASES}
            site = mapped["pos_site_code"]
            biz_date = _to_iso_date(mapped["business_date"])
            if site is None or str(site).strip() == "" or biz_date is None:
                skipped.append({"row": {k: self._jsonable(v) for k, v in raw.items()},
                                "reason": "missing or unparseable site/date"})
                continue
            site = str(site).strip()

            key = (site, biz_date)
            bucket = buckets.setdefault(key, {
                "pos_site_code": site, "business_date": biz_date,
                "gross_sales": None, "vat": None, "discounts": None,
                "net_sales": None, "covers": None, "lines": [],
            })

            is_item_row = mapped["pos_item_code"] is not None or (
                mapped["item_name"] is not None and mapped["qty"] is not None
            )
            if is_item_row:
                qty = _to_decimal(mapped["qty"])
                gross_value = _to_decimal(mapped["gross_value"])
                bucket["lines"].append({
                    "pos_item_code": str(mapped["pos_item_code"]).strip() if mapped["pos_item_code"] is not None else None,
                    "item_name": str(mapped["item_name"]).strip() if mapped["item_name"] is not None else None,
                    "qty": float(qty) if qty is not None else None,
                    "gross_value": float(gross_value) if gross_value is not None else None,
                })
                # Item rows may also carry the summary columns; accumulate them.
                self._accumulate(bucket, "gross_sales", _to_decimal(mapped["gross_sales"]) or gross_value)
                self._accumulate(bucket, "vat", _to_decimal(mapped["vat"]))
                self._accumulate(bucket, "discounts", _to_decimal(mapped["discounts"]))
                self._accumulate(bucket, "net_sales", _to_decimal(mapped["net_sales"]))
                self._accumulate(bucket, "covers", _to_decimal(mapped["covers"]))
            else:
                # Summary row for the site/day — accumulate (multiple rows, e.g.
                # per revenue centre, sum up).
                gross = _to_decimal(mapped["gross_sales"])
                if gross is None and mapped["net_sales"] is None:
                    skipped.append({"row": {k: self._jsonable(v) for k, v in raw.items()},
                                    "reason": "no sales value columns recognised"})
                    continue
                self._accumulate(bucket, "gross_sales", gross)
                self._accumulate(bucket, "vat", _to_decimal(mapped["vat"]))
                self._accumulate(bucket, "discounts", _to_decimal(mapped["discounts"]))
                self._accumulate(bucket, "net_sales", _to_decimal(mapped["net_sales"]))
                self._accumulate(bucket, "covers", _to_decimal(mapped["covers"]))

        out: list[dict] = []
        for bucket in buckets.values():
            gross = bucket["gross_sales"]
            vat = bucket["vat"]
            discounts = bucket["discounts"]
            net = bucket["net_sales"]
            if net is None and gross is not None:
                net = gross - (vat or Decimal(0)) - (discounts or Decimal(0))
            row = {
                "pos_site_code": bucket["pos_site_code"],
                "business_date": bucket["business_date"],
                "gross_sales": float(gross) if gross is not None else None,
                "vat": float(vat) if vat is not None else None,
                "discounts": float(discounts) if discounts is not None else None,
                "net_sales": float(net) if net is not None else None,
                "covers": int(bucket["covers"]) if bucket["covers"] is not None else None,
            }
            if bucket["lines"]:
                row["lines"] = bucket["lines"]
            out.append(row)

        payload["skipped_rows"] = skipped
        return out

    @staticmethod
    def _accumulate(bucket: dict, field: str, value: Decimal | None) -> None:
        if value is None:
            return
        bucket[field] = value if bucket[field] is None else bucket[field] + value

    @staticmethod
    def _pick(raw: dict, field: str):
        for alias in HEADER_ALIASES[field]:
            if alias in raw and raw[alias] is not None and str(raw[alias]).strip() != "":
                return raw[alias]
        return None

    @staticmethod
    def _jsonable(v):
        if isinstance(v, (datetime, date)):
            return v.isoformat()
        if isinstance(v, Decimal):
            return float(v)
        return v
